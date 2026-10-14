#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""SAP XEP TAI LIEU THEO NAM VA DAT TEN THEO QUY DINH 68-QD/TW
================================================================

Cong cu nay se:

1. Quet toan bo file .pdf / .docx / .doc trong thu muc nguon (ke ca thu muc con).
2. Doc noi dung tung file de tim NAM (va ngay, tac gia neu co).
3. Copy file vao thu muc dich, chia theo nam:  <thu muc dich>\\<nam>\\
4. Doi ten file theo cau truc:  <TenChinh><DDmmYY><Loai><DonVi><M_T><SoLanXB>.<ext>
   (xem tools/naming.py de biet y nghia tung nhom).
5. Ghi 1 file Excel bao cao gom:
     - Da_xu_ly                    : moi file da xu ly (ten cu, ten moi, nam, tac gia...)
     - Thieu_nam_trong_noi_dung    : file khong doc duoc nam (kem ngay chinh sua file)
     - Thieu_ngay_trong_noi_dung   : file khong doc duoc ngay trong noi dung
     - Loi_doc_file                : file loi / khong doc duoc noi dung
     - Khong_ho_tro                : file khac dinh dang pdf/docx/doc
     - Tong_hop_theo_nam           : thong ke theo nam
     - Cau_hinh                    : tham so da dung khi chay

Vi du chay:
    python sap_xep_tai_lieu.py                 # dung config.json
    python sap_xep_tai_lieu.py --dry-run       # chi xem truoc, khong copy
    python sap_xep_tai_lieu.py --limit 5       # thu voi 5 file dau
"""
import argparse
import datetime as dt
import json
import os
import shutil
import sys
import time

if getattr(sys, "frozen", False):
    HERE = os.path.dirname(sys.executable)          # thư mục chứa file .exe
else:
    HERE = os.path.dirname(os.path.abspath(__file__))
    
TOOLS = os.path.join(HERE, "tools")
LIBS = os.path.join(HERE, "_libs")
for p in (TOOLS, LIBS, HERE):
    if os.path.isdir(p) and p not in sys.path:
        sys.path.insert(0, p)

from text_extract import extract_text, ExtractError  # noqa: E402
from meta_extract import analyze, ten_sach_ung_vien  # noqa: E402
import naming  # noqa: E402
import bang_ke  # noqa: E402

MAC_DINH_CFG = "config.json"


# --------------------------------------------------------------------------- #
# Bang ke sach xuat ban (tra Nhom 4 theo cot BBT)
# --------------------------------------------------------------------------- #
def nap_bang_ke(cfg):
    """Nap bang ke sach (Tong hop SXB ...xlsx). Tra ve BangKeSach hoac None."""
    c = cfg.get("bang_ke_sach", {})
    if not c.get("bat", True):
        print("Bang ke sach  : (tat trong config)")
        return None
    duong_dan = (c.get("file") or "").strip()
    if not duong_dan:
        return None
    if not os.path.isabs(duong_dan):
        duong_dan = os.path.join(HERE, duong_dan)
    if not os.path.isfile(duong_dan):
        print("!! Khong thay file bang ke sach: %s" % duong_dan)
        print("   -> Nhom 4 se lay tu ten file cu / gia tri mac dinh.")
        return None
    try:
        bk = bang_ke.BangKeSach(duong_dan,
                                nguong=c.get("nguong_khop", 0.6),
                                tim_nam_lan_can=c.get("tim_nam_lan_can", 1),
                                ban_do=cfg.get("ban_bien_tap", {}))
        bk._mo()
        print("Bang ke sach  : %s (%d sheet nam: %s..%s)" % (
            os.path.basename(duong_dan), len(bk.ten_sheet),
            bk.ten_sheet[0] if bk.ten_sheet else "?", bk.ten_sheet[-1] if bk.ten_sheet else "?"))
        return bk
    except Exception as e:
        print("!! Loi doc bang ke sach: %s: %s" % (type(e).__name__, e))
        return None


# --------------------------------------------------------------------------- #
# Cau hinh
# --------------------------------------------------------------------------- #
def load_config(path):
    # utf-8-sig de doc duoc ca file config co BOM (do Notepad/PowerShell luu)
    with open(path, encoding="utf-8-sig") as fh:
        cfg = json.load(fh)
    cfg.setdefault("ten", {})
    cfg.setdefault("mac_dinh", {})
    return cfg


def quet_file(src, cfg):
    """Liet ke file can xu ly (duyet ca thu muc con)."""
    ho_tro = {e.lower() for e in cfg.get("dinh_dang_ho_tro", [".pdf", ".docx", ".doc"])}
    dst = os.path.abspath(cfg.get("thu_muc_dich", ""))
    bo_qua_tam = cfg.get("bo_qua_file_tam", True)
    ds, bo_qua = [], []
    for root, dirs, files in os.walk(src):
        dirs[:] = [d for d in dirs if os.path.abspath(os.path.join(root, d)) != dst]
        for fn in sorted(files):
            p = os.path.join(root, fn)
            ext = os.path.splitext(fn)[1].lower()
            if bo_qua_tam and fn.startswith("~$"):
                bo_qua.append((p, "file tam cua Word (~$) - thuong la file rac"))
                continue
            if ext not in ho_tro:
                bo_qua.append((p, "dinh dang %s khong nam trong danh sach ho tro" % (ext or "(khong co)")))
                continue
            ds.append(p)
    ds.sort()
    return ds, bo_qua


# --------------------------------------------------------------------------- #
# Chon ngay (nhom 2)
# --------------------------------------------------------------------------- #
def chon_ngay(info, old, mtime, cfg, nam_thu_muc=None):
    """Tra ve (dd, mm, yy, nguon, canh_bao, khoa_nguon).

    Thu tu:
      1. Ung vien co nam trung voi nam thu muc (theo thu tu uu tien trong config).
      2. Ngay co san trong ten file cu (ngay tao file that) - giu nguyen nam,
         chi canh bao neu khac nam thu muc.
      3. Ung vien con lai, ep nam theo thu muc (co canh bao).
    """
    ung_vien = {}

    if old.get("ngay_hop_le"):
        n = old["ngay"]
        ung_vien["ten_file_cu"] = (int(n[:2]), int(n[2:4]), int(n[4:]),
                                   "6 so trong ten file cu (%s)" % n)

    d = info.get("ngay_du")
    if d:
        ung_vien["noi_dung_du_ngay"] = (d["ngay"], d["thang"], d["nam"] % 100,
                                        "noi dung: %s" % d["chuoi"][:40])

    m = info.get("ngay_thang")
    if m:
        ung_vien["noi_dung_thang_nam"] = (1, m["thang"], m["nam"] % 100,
                                          "noi dung chi co thang/nam: %s" % m["chuoi"][:40])

    if mtime:
        ung_vien["ngay_sua_file"] = (mtime.day, mtime.month, mtime.year % 100,
                                     "ngay chinh sua file tren o dia")

    def _cb(key):
        cb = []
        if key == "noi_dung_thang_nam":
            cb.append("noi dung chi co thang/nam, da lay ngay = 01")
        if key == "ngay_sua_file":
            cb.append("khong doc duoc ngay tao file, da dung ngay chinh sua file tren o dia")
        return cb

    def _cung_nam(yy):
        return (nam_thu_muc is None) or (yy + 2000 == nam_thu_muc) or (yy + 1900 == nam_thu_muc)

    thu_tu = [k for k in cfg.get("ngay_uu_tien", []) if k in ung_vien]
    phai_cung_nam = bool(cfg.get("ngay_phai_cung_nam_voi_thu_muc", True))

    # 1. uu tien ung vien cung nam voi thu muc
    if phai_cung_nam and nam_thu_muc:
        for key in thu_tu:
            dd, mm, yy, nguon = ung_vien[key]
            if _cung_nam(yy):
                return dd, mm, yy, nguon, _cb(key), key

        # 2. ngay tao file that trong ten cu -> giu nguyen, chi canh bao
        if "ten_file_cu" in ung_vien:
            dd, mm, yy, nguon = ung_vien["ten_file_cu"]
            return dd, mm, yy, nguon, [
                "ngay tao file (%02d/%02d/20%02d) khac nam xuat ban (%s) cua thu muc"
                % (dd, mm, yy, nam_thu_muc)], "ten_file_cu"

        # 3. khong con cach nao khac: ep nam theo thu muc
        if thu_tu:
            dd, mm, yy, nguon = ung_vien[thu_tu[0]]
            return dd, mm, nam_thu_muc % 100, nguon, [
                "%s khong cung nam voi thu muc (%s), da doi nam trong ten file thanh %02d"
                % (nguon, nam_thu_muc, nam_thu_muc % 100)], thu_tu[0]

    for key in thu_tu:
        dd, mm, yy, nguon = ung_vien[key]
        return dd, mm, yy, nguon, _cb(key), key
    return None, None, None, "", [], None


# --------------------------------------------------------------------------- #
# Xu ly 1 file
# --------------------------------------------------------------------------- #
def xu_ly_file(path, src, cfg):
    ten_goc = os.path.basename(path)
    ext = os.path.splitext(ten_goc)[1].lower()
    st = os.stat(path)
    tm_goc = os.path.relpath(os.path.dirname(path), src)
    rec = dict(
        ten_goc=ten_goc, duong_dan_goc=path,
        thu_muc_goc="(ngay tai thu muc nguon)" if tm_goc == "." else tm_goc,
        ext=ext,
        ngay_sua=dt.datetime.fromtimestamp(st.st_mtime),
        kich_thuoc_kb=round(st.st_size / 1024.0, 1),
        trang_thai="ok", canh_bao=[],
        ten_moi="", thu_muc_dich="",
        nam=None, nam_cach="", nam_tin_cay="", nam_bang_chung="",
        tac_gia="", tac_gia_cach="",
        ngay_ten="", nguon_ngay="", ngay_noi_dung="", ngay_ten_file_cu="",
        ngay_sua_ngan="", ten_chinh="", loai="", don_vi="", tinh_trang="",
        so_lan="", nguon_metadata="", nguon_nhom4="",
        bk_bbt="", bk_ma="", bk_ten="", bk_tac_gia="", bk_vi_tri="", bk_diem="", bk_lan_xb="",
        bk_nguon_ten="",
    )

    old = naming.parse_old_name(ten_goc)
    if old.get("ngay_hop_le"):
        n = old["ngay"]
        rec["ngay_ten_file_cu"] = "%s/%s/%s" % (n[:2], n[2:4], n[4:])

    # 1. doc noi dung
    try:
        text, _dec = extract_text(path, max_pages=int(cfg.get("so_trang_pdf_toi_da")) or None)
    except ExtractError as e:
        rec["trang_thai"] = "loi_doc"
        rec["loi"] = str(e)
        return rec
    except Exception as e:  # loi khong luong truoc
        rec["trang_thai"] = "loi_doc"
        rec["loi"] = "%s: %s" % (type(e).__name__, e)
        return rec

    info = analyze(text)

    # 2. nam -> thu muc
    rec["nam"] = info["nam"]
    rec["nam_cach"] = info["nam_cach"]
    rec["nam_tin_cay"] = info["nam_tin_cay"]
    rec["nam_bang_chung"] = info["nam_bang_chung"]
    rec["tac_gia"] = info["tac_gia"]
    rec["tac_gia_cach"] = info["tac_gia_cach"]
    d = info.get("ngay_du") or info.get("ngay_thang")
    if d:
        rec["ngay_noi_dung"] = "%02d/%02d/%04d" % (d["ngay"], d["thang"], d["nam"])

    if not info["nam"]:
        rec["trang_thai"] = "thieu_nam"
        rec["canh_bao"].append("khong doc duoc nam trong noi dung")
        ten_moi = ten_goc  # giu nguyen ten, cho nguoi dung xu ly tay
        rec["ten_moi"] = ten_moi
        rec["thu_muc_dich"] = os.path.join(cfg["thu_muc_dich"],
                                           cfg.get("thu_muc_khong_xac_dinh_nam", "_KHONG_XAC_DINH_NAM"))
        rec["ten_chinh"] = naming.normalize_base(os.path.splitext(ten_goc)[0], cfg)
        return rec

    # 3. ngay (nhom 2)
    dd, mm, yy, nguon_ngay, cb, key_ngay = chon_ngay(info, old, rec["ngay_sua"], cfg, info["nam"])
    rec["canh_bao"].extend(cb)
    if dd is None:
        dd, mm, yy = 1, 1, info["nam"] % 100
        nguon_ngay = "khong tim thay ngay nao -> dung 01/01/%02d" % yy
        rec["canh_bao"].append("khong co ngay trong noi dung va ten file, da dung 01/01")
        key_ngay = "khong_co"
    rec["ngay_ten"] = "%02d%02d%02d" % (dd, mm, yy)
    rec["nguon_ngay"] = nguon_ngay
    rec["ngay_sua_ngan"] = rec["ngay_sua"].strftime("%d/%m/%Y")

    # 4. ten chinh (nhom 1) - dung de DAT TEN, va ten sach THO de TRA BANG KE
    base_goc = old["base"] if old.get("co_duoi") else os.path.splitext(ten_goc)[0]
    rec["ten_chinh"] = naming.normalize_base(base_goc, cfg)

    # 5. tra bang ke sach:  ten sach + nam  ->  cot BBT  ->  ma Nhom 4
    #    Tra bang TEN SACH THO (chua chuan hoa): ten file bo nhom ma cu, roi den
    #    ten file nguyen ban, roi cac dong co the la ten sach o dau tai lieu.
    #    Viec bo dau / bo ky tu dac biet do bang_ke.chuan_hoa() lam cho CA HAI ve.
    kq_bk = None
    nguon_ten_bk = ""
    bk = cfg.get("_bang_ke")
    if bk is not None:
        c_bk = cfg.get("bang_ke_sach", {})
        nguong_nd = float(c_bk.get("nguong_khop_noi_dung", 0.8))
        cho = [(0, "ten file", base_goc)]
        if base_goc != rec["ten_chinh"]:
            cho.append((0, "ten file (da chuan hoa)", rec["ten_chinh"]))
        for j, s in enumerate(ten_sach_ung_vien(text), 1):
            cho.append((1, "noi dung (dong %d)" % j, s))

        de_xuat = None                      # (diem, kq, nguon) - dung khi khong dat nguong
        for uu_tien, nguon_ten, ten in cho:
            try:
                kq = bk.tra(ten, info["nam"], info["tac_gia"])
            except Exception as e:
                rec["canh_bao"].append("loi tra bang ke sach: %s: %s" % (type(e).__name__, e))
                continue
            if not kq:
                continue
            if uu_tien == 0 and not kq["dat_nguong"] and c_bk.get("tim_ngoai_nam", True):
                try:
                    kq2 = bk.tra(ten, info["nam"], info["tac_gia"], mo_rong=True)
                    if kq2 and kq2["dat_nguong"]:
                        kq = kq2
                except Exception:
                    pass
            if de_xuat is None or kq["diem"] > de_xuat[0]:
                de_xuat = (kq["diem"], kq, nguon_ten)
            if not kq["dat_nguong"]:
                continue
            if uu_tien == 1 and kq["diem"] < nguong_nd:
                continue                    # dong trong noi dung phai khop ro moi dung
            kq_bk, nguon_ten_bk = kq, nguon_ten
            break
        if kq_bk is None and de_xuat:
            # chi dung de BAO CAO (goi y dong khop nhat), khong dung de dat ten
            kq_bk, nguon_ten_bk = de_xuat[1], de_xuat[2]
            kq_bk["dat_nguong"] = False
    if kq_bk:
        rec["bk_bbt"] = kq_bk["bbt"]
        rec["bk_ma"] = kq_bk["ma"] or ""
        rec["bk_ten"] = kq_bk["ten_sach"]
        rec["bk_tac_gia"] = kq_bk["tac_gia"]
        rec["bk_nguon_ten"] = nguon_ten_bk
        rec["bk_vi_tri"] = "sheet %s, dong %s" % (kq_bk["sheet"], kq_bk["dong"])
        rec["bk_diem"] = kq_bk["diem"]
        rec["bk_lan_xb"] = kq_bk["lan_xb"]
        if not kq_bk["dat_nguong"]:
            rec["canh_bao"].append(
                "ten sach chi khop %.2f voi bang ke (%s: '%s') - KHONG dung ma nay, can kiem tra lai"
                % (kq_bk["diem"], rec["bk_vi_tri"], kq_bk["ten_sach"][:50]))
        else:
            if kq_bk.get("ngoai_nam"):
                rec["canh_bao"].append(
                    "nam %s khong co trong bang ke; da lay theo sach cung ten o sheet %s (diem %.2f) - can kiem tra lai"
                    % (info["nam"], kq_bk["sheet"], kq_bk["diem"]))
            elif kq_bk["lech_nam"]:
                rec["canh_bao"].append("sach nam trong sheet %s, lech %d nam so voi nam xuat ban"
                                       % (kq_bk["sheet"], kq_bk["lech_nam"]))
            if not kq_bk["chac_chan"]:
                rec["canh_bao"].append(
                    "khop ten sach khong chac chan (diem %.2f, muc 2 la %.2f)"
                    % (kq_bk["diem"], kq_bk["diem_2"]))
            if kq_bk["ghi_chu_ma"]:
                rec["canh_bao"].append(kq_bk["ghi_chu_ma"])
    elif bk is not None:
        rec["canh_bao"].append("khong tim thay ten sach trong bang ke nam %s" % info["nam"])

    # tac gia: neu noi dung khong doc duoc thi lay tu bang ke (chi khi khop dat nguong)
    if not rec["tac_gia"] and kq_bk and kq_bk["dat_nguong"] and kq_bk.get("tac_gia"):
        rec["tac_gia"] = kq_bk["tac_gia"]
        rec["tac_gia_cach"] = "bang ke sach (%s)" % rec["bk_vi_tri"]

    # 6. cac nhom 3-6
    nguon_md = []
    if old.get("loai"):
        rec["loai"], n1 = old["loai"], "ten file cu"
    else:
        rec["loai"], n1 = naming.guess_loai_file(text, cfg, cfg["mac_dinh"].get("loai_file", "R"))
        rec["canh_bao"].append("nhom 3 (loai file) phai doan: %s" % n1)

    # --- Nhom 4: bang ke sach truoc, roi den ten file cu, cuoi cung la mac dinh
    ma_bk = (kq_bk or {}).get("ma") if (kq_bk and kq_bk["dat_nguong"]) else None
    if ma_bk:
        rec["don_vi"], n2 = ma_bk, "bang ke sach (BBT=%s -> %s)" % (kq_bk["bbt"], ma_bk)
        if old.get("don_vi") and old["don_vi"] != ma_bk:
            rec["canh_bao"].append("nhom 4: bang ke ghi %s nhung ten file cu ghi %s -> da lay theo bang ke"
                                   % (ma_bk, old["don_vi"]))
    elif old.get("don_vi"):
        rec["don_vi"], n2 = old["don_vi"], "ten file cu"
        if kq_bk and kq_bk["bbt"]:
            rec["canh_bao"].append("bang ke co BBT '%s' nhung chua map duoc ma -> tam dung ma trong ten file cu (%s)"
                                   % (kq_bk["bbt"], old["don_vi"]))
    else:
        rec["don_vi"] = cfg["mac_dinh"].get("don_vi", "04")
        n2 = "mac dinh config"
        rec["canh_bao"].append("nhom 4: khong tra duoc bang ke va ten file cu khong co ma -> dung mac dinh %s"
                               % rec["don_vi"])
    rec["nguon_nhom4"] = n2

    if old.get("tinh_trang") and old.get("so_lan") is not None:
        rec["tinh_trang"], rec["so_lan"] = old["tinh_trang"], old["so_lan"]
        n3 = "ten file cu"
    else:
        tt, sl, n3 = naming.guess_tinh_trang(text, cfg, cfg["mac_dinh"].get("tinh_trang", "M"))
        rec["tinh_trang"] = old.get("tinh_trang") or tt
        rec["so_lan"] = old["so_lan"] if old.get("so_lan") is not None else sl
        rec["canh_bao"].append("nhom 5-6 (M/T, so lan XB) phai doan: %s" % n3)
    nguon_md = [n1, n2, n3]
    rec["nguon_metadata"] = " | ".join(nguon_md)

    # 7. ten moi
    rec["ten_moi"] = naming.build_name(rec["ten_chinh"], rec["ngay_ten"], rec["loai"],
                                       rec["don_vi"], rec["tinh_trang"], rec["so_lan"],
                                       ext, cfg)
    rec["thu_muc_dich"] = os.path.join(cfg["thu_muc_dich"], str(info["nam"]))

    # 8. canh bao chat luong
    if rec["nam_tin_cay"] == "thap":
        rec["canh_bao"].append("nam doc duoc khong chac chan (%s)" % rec["nam_cach"])
    if yy + 2000 != info["nam"] and yy + 1900 != info["nam"]:
        rec["canh_bao"].append("nam trong ten file (%02d) khac nam thu muc (%s)" % (yy, info["nam"]))
    if old.get("ngay_hop_le") and old["ngay"] != rec["ngay_ten"]:
        rec["canh_bao"].append("ngay trong ten file cu (%s) khac ngay da dung (%s)"
                               % (old["ngay"], rec["ngay_ten"]))
    if old.get("ngay_hop_le"):
        y_cu = int(old["ngay"][4:]) + (2000 if int(old["ngay"][4:]) < 80 else 1900)
        if y_cu != info["nam"]:
            rec["canh_bao"].append("nam trong ten file cu (20%02d) khac nam doc tu noi dung (%s)"
                                   % (int(old["ngay"][4:]), info["nam"]))
    return rec


# --------------------------------------------------------------------------- #
# Ghi bao cao Excel
# --------------------------------------------------------------------------- #
def ghi_excel(rows, loi, bo_qua, cfg, nguon, thoi_diem, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    wrap = Alignment(vertical="top", wrap_text=True)

    def tao_sheet(ws, tieu_de, cols, data):
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in data:
            ws.append(r)
        for c in range(1, len(cols) + 1):
            ws.column_dimensions[get_column_letter(c)].width = tieu_de[c - 1]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), ws.max_row)

    # ---- Da_xu_ly
    ws = wb.active
    ws.title = "Da_xu_ly"
    cols = ["STT", "Ten file goc", "Ten file moi", "Thu muc (nam)", "Nam", "Nguon nam",
            "Do tin cay nam", "Ten chinh (nhom 1)", "Ngay (nhom 2)", "Nguon ngay",
            "Loai (N3)", "Don vi (N4)", "Nguon nhom 4", "BBT (bang ke)", "Ma N4 (bang ke)",
            "Ten sach trong bang ke", "Ten sach lay tu", "Vi tri bang ke", "Diem khop",
            "Tac gia (bang ke)", "Lan XB (bang ke)", "Tinh trang (N5)", "So lan XB (N6)",
            "Nguon metadata", "Tac gia", "Cach xac dinh tac gia", "Ngay noi dung",
            "Ngay trong ten file cu", "Ngay sua file", "Kich thuoc (KB)",
            "Thu muc goc", "Duong dan goc", "Duong dan moi", "Canh bao"]
    widths = [5, 46, 46, 12, 7, 26, 12, 40, 10, 30, 8, 8, 30, 12, 12, 42, 18, 18, 10, 26, 12,
              10, 10, 22, 26, 24, 12, 14, 12, 11, 30, 50, 50, 44]
    data = []
    for i, r in enumerate(rows, 1):
        data.append([i, r["ten_goc"], r["ten_moi"], str(r["nam"]), r["nam"],
                      r["nam_cach"], r["nam_tin_cay"], r["ten_chinh"], r["ngay_ten"],
                      r["nguon_ngay"], r["loai"], r["don_vi"], r.get("nguon_nhom4", ""),
                      r.get("bk_bbt", ""), r.get("bk_ma", ""), r.get("bk_ten", ""),
                      r.get("bk_nguon_ten", ""), r.get("bk_vi_tri", ""), r.get("bk_diem", ""),
                      r.get("bk_tac_gia", ""), r.get("bk_lan_xb", ""), r["tinh_trang"], r["so_lan"],
                      r["nguon_metadata"], r["tac_gia"], r["tac_gia_cach"],
                      r["ngay_noi_dung"], r["ngay_ten_file_cu"],
                      r["ngay_sua"].strftime("%d/%m/%Y %H:%M"), r["kich_thuoc_kb"],
                      r["thu_muc_goc"], r["duong_dan_goc"],
                      os.path.join(r["thu_muc_dich"], r["ten_moi"]) if r["ten_moi"] else "",
                      " ; ".join(r["canh_bao"])])
    tao_sheet(ws, widths, cols, data)

    # ---- Thieu_nam_trong_noi_dung
    ws2 = wb.create_sheet("Thieu_nam_trong_noi_dung")
    cols2 = ["STT", "Ten file", "Duong dan", "Ngay chinh sua file", "Kich thuoc (KB)",
             "Nam doc duoc tu ten file cu", "Ghi chu"]
    widths2 = [5, 50, 60, 20, 12, 22, 50]
    data2 = []
    for i, r in enumerate([x for x in loi if x["trang_thai"] == "thieu_nam"], 1):
        nam_cu = ""
        o = naming.parse_old_name(r["ten_goc"])
        if o.get("ngay_hop_le"):
            nam_cu = "20%s" % o["ngay"][4:]
        data2.append([i, r["ten_goc"], r["duong_dan_goc"],
                      r["ngay_sua"].strftime("%d/%m/%Y %H:%M"), r["kich_thuoc_kb"],
                      nam_cu, " ; ".join(r["canh_bao"])])
    tao_sheet(ws2, widths2, cols2, data2)

    # ---- Thieu_ngay_trong_noi_dung
    ws3 = wb.create_sheet("Thieu_ngay_trong_noi_dung")
    cols3 = ["STT", "Ten file goc", "Duong dan", "Ngay chinh sua file",
             "Ngay da dung (nhom 2)", "Nguon ngay da dung", "Nam", "Ghi chu"]
    widths3 = [5, 50, 60, 20, 16, 34, 8, 50]
    data3 = []
    for i, r in enumerate([x for x in rows if not x["ngay_noi_dung"]], 1):
        data3.append([i, r["ten_goc"], r["duong_dan_goc"],
                      r["ngay_sua"].strftime("%d/%m/%Y %H:%M"), r["ngay_ten"],
                      r["nguon_ngay"], r["nam"], " ; ".join(r["canh_bao"])])
    tao_sheet(ws3, widths3, cols3, data3)

    # ---- Loi_doc_file
    ws4 = wb.create_sheet("Loi_doc_file")
    cols4 = ["STT", "Ten file", "Duong dan", "Ngay chinh sua file", "Kich thuoc (KB)",
             "Loai loi", "Chi tiet loi"]
    widths4 = [5, 50, 60, 20, 12, 16, 90]
    data4 = []
    for i, r in enumerate([x for x in loi if x["trang_thai"] == "loi_doc"], 1):
        data4.append([i, r["ten_goc"], r["duong_dan_goc"],
                      r["ngay_sua"].strftime("%d/%m/%Y %H:%M"), r["kich_thuoc_kb"],
                      "khong doc duoc", r.get("loi", "")])
    tao_sheet(ws4, widths4, cols4, data4)

    # ---- Khong_ho_tro
    ws5 = wb.create_sheet("Khong_ho_tro")
    cols5 = ["STT", "Ten file", "Duong dan", "Ly do"]
    widths5 = [5, 55, 70, 60]
    data5 = [[i, os.path.basename(p), p, ly] for i, (p, ly) in enumerate(bo_qua, 1)]
    tao_sheet(ws5, widths5, cols5, data5)

    # ---- Tong_hop_theo_nam
    ws6 = wb.create_sheet("Tong_hop_theo_nam")
    cols6 = ["Nam", "So file", "Dung luong (MB)", "So file do tin cay thap",
             "So file lech nam voi ten cu", "Thu muc"]
    widths6 = [8, 10, 16, 20, 24, 40]
    theo_nam = {}
    for r in rows:
        k = r["nam"]
        t = theo_nam.setdefault(k, dict(n=0, kb=0.0, low=0, lech=0))
        t["n"] += 1
        t["kb"] += r["kich_thuoc_kb"]
        if r["nam_tin_cay"] == "thap":
            t["low"] += 1
        if any("khac nam xuat ban" in c or "khong cung nam voi thu muc" in c
               for c in r["canh_bao"]):
            t["lech"] += 1
    data6 = []
    for k in sorted(theo_nam, key=lambda x: (x is None, x)):
        t = theo_nam[k]
        data6.append([k if k else "(khong xac dinh)", t["n"], round(t["kb"] / 1024.0, 2),
                      t["low"], t["lech"],
                      os.path.join(cfg["thu_muc_dich"], str(k)) if k else
                      os.path.join(cfg["thu_muc_dich"], cfg.get("thu_muc_khong_xac_dinh_nam", ""))])
    tao_sheet(ws6, widths6, cols6, data6)

    # ---- Cau_hinh
    ws7 = wb.create_sheet("Cau_hinh")
    ws7.append(["Tham so", "Gia tri"])
    ws7["A1"].font = head_font
    ws7["A1"].fill = head_fill
    ws7["B1"].font = head_font
    ws7["B1"].fill = head_fill
    ws7.column_dimensions["A"].width = 34
    ws7.column_dimensions["B"].width = 100
    for k in ("thu_muc_nguon", "thu_muc_dich", "dinh_dang_ho_tro", "ngay_uu_tien",
              "copy_giu_file_goc", "thu_muc_khong_xac_dinh_nam"):
        ws7.append([k, json.dumps(cfg.get(k), ensure_ascii=False)])
    ws7.append(["thoi_diem_chay", thoi_diem.strftime("%d/%m/%Y %H:%M:%S")])
    ws7.append(["tong so file xu ly", len(rows)])
    ws7.append(["tong so file loi", len([x for x in loi if x["trang_thai"] == "loi_doc"])])
    ws7.append(["tong so file thieu nam", len([x for x in loi if x["trang_thai"] == "thieu_nam"])])
    ws7.append(["tong so file bo qua", len(bo_qua)])
    bk = cfg.get("_bang_ke")
    if bk is not None:
        ws7.append(["bang ke sach (file)", os.path.basename(bk.duong_dan)])
        ws7.append(["bang ke sach (nguong khop)", bk.nguong])
        ws7.append(["so file tra duoc bang ke",
                    len([x for x in rows if x.get("bk_diem") not in ("", None)])])
        ws7.append(["so file lay duoc ma Nhom 4 tu bang ke",
                    len([x for x in rows if str(x.get("nguon_nhom4", "")).startswith("bang ke")])])
    for row in ws7.iter_rows(min_row=2):
        row[1].alignment = wrap

    # ---- BBT_chua_map: cac ma BBT trong bang ke chua co trong config
    ws8 = wb.create_sheet("BBT_chua_map")
    cols8 = ["STT", "Ma BBT trong bang ke", "So lan gap", "Vi du ten sach",
             "Goi y: bo sung vao config.json"]
    widths8 = [5, 24, 12, 60, 46]
    data8 = []
    bk = cfg.get("_bang_ke")
    if bk is not None:
        for i, (ma, (so_lan, vi_du)) in enumerate(
                sorted(bk.chua_map.items(), key=lambda kv: -kv[1][0]), 1):
            data8.append([i, ma, so_lan, vi_du,
                          '"%s": "<so nhom 4 01..21>"' % bang_ke.chuan_hoa(ma)])
    if not data8:
        data8.append(["", "(khong co)", "", "Tat ca ma BBT trong bang ke deu da map duoc.", ""])
    tao_sheet(ws8, widths8, cols8, data8)

    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def kiem_tra_thu_vien():
    """Bao loi ro rang neu thieu thu vien can thiet."""
    thieu = []
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        thieu.append("openpyxl (ghi bao cao Excel)")
    try:
        import pypdf  # noqa: F401
    except ImportError:
        thieu.append("pypdf (doc file PDF)")
    if thieu:
        print("!! Thieu thu vien: %s" % ", ".join(thieu))
        print("   Cach xu ly:  python tools\\download_libs.py")
        print("   hoac:        pip install pypdf openpyxl")
        return False
    return True

def _hoi(cau_hoi):
    return input(cau_hoi).strip().strip('"').strip()

def main():
    ap = argparse.ArgumentParser(description="Sap xep tai lieu theo nam va dat ten theo QD 68-QD/TW")
    ap.add_argument("--config", default=os.path.join(HERE, MAC_DINH_CFG), help="file cau hinh JSON")
    ap.add_argument("--input", help="thu muc nguon (ghi de config)")
    ap.add_argument("--output", help="thu muc dich (ghi de config)")
    ap.add_argument("--dry-run", action="store_true", help="chi xem truoc, khong copy file")
    ap.add_argument("--move", action="store_true", help="di chuyen file thay vi copy")
    ap.add_argument("--limit", type=int, default=0, help="chi xu ly N file dau (de thu)")
    ap.add_argument("-v", "--verbose", action="store_true",
                    help="in ra tung file: ten cu -> ten moi, nam, tac gia")
    ap.add_argument("--no-bang-ke", action="store_true",
                    help="khong tra bang ke sach (Nhom 4 lay tu ten file cu / mac dinh)")
    args = ap.parse_args()

    if not kiem_tra_thu_vien():
        return 3

    cfg = load_config(args.config)
    
    if not cfg.get("thu_muc_nguon"):
        cfg["thu_muc_nguon"] = _hoi("Kéo thả thư mục nguồn (thư mục chứa file cần sắp xếp) vào đây rồi nhấn Enter: ")
    if not cfg.get("thu_muc_dich"):
        cfg["thu_muc_dich"] = _hoi("Kéo thả thư mục đích (thư mục để chứa file đã sắp xếp) vào đây rồi nhấn Enter: ")

    if args.input:
        cfg["thu_muc_nguon"] = args.input
    if args.output:
        cfg["thu_muc_dich"] = args.output
    if args.move:
        cfg["copy_giu_file_goc"] = False
    if args.limit:
        cfg["gioi_han_so_file"] = args.limit

    src = os.path.abspath(cfg["thu_muc_nguon"])
    dst = os.path.abspath(cfg["thu_muc_dich"])
    if not os.path.isdir(src):
        print("!! Khong ton tai thu muc nguon: %s" % src)
        return 2
    if args.dry_run:
        print(">>> CHE DO XEM TRUOC (dry-run): khong copy file nao.\n")

    ds_file, bo_qua = quet_file(src, cfg)
    if cfg.get("gioi_han_so_file"):
        ds_file = ds_file[:int(cfg["gioi_han_so_file"])]

    print("Thu muc nguon : %s" % src)
    print("Thu muc dich  : %s" % dst)
    print("So file can xu ly: %d   |   bo qua: %d\n" % (len(ds_file), len(bo_qua)))

    # nap bang ke sach (tra Nhom 4 theo cot BBT)
    if not args.no_bang_ke:
        cfg["_bang_ke"] = nap_bang_ke(cfg)
    else:
        print("Bang ke sach  : (tat bang --no-bang-ke)")
        cfg["_bang_ke"] = None
    print()

    t0 = time.time()
    rows, loi = [], []
    da_dung = set()
    for i, p in enumerate(ds_file, 1):
        rel = os.path.relpath(p, src)
        try:
            rec = xu_ly_file(p, src, cfg)
        except Exception as e:
            rec = dict(ten_goc=os.path.basename(p), duong_dan_goc=p, thu_muc_goc="",
                       ext=os.path.splitext(p)[1], kich_thuoc_kb=0,
                       ngay_sua=dt.datetime.fromtimestamp(os.path.getmtime(p)),
                       trang_thai="loi_doc", loi="%s: %s" % (type(e).__name__, e),
                       canh_bao=[], ten_moi="", thu_muc_dich="", nam=None, nam_cach="",
                       nam_tin_cay="", nam_bang_chung="", tac_gia="", tac_gia_cach="",
                       ngay_ten="", nguon_ngay="", ngay_noi_dung="", ngay_ten_file_cu="",
                       ngay_sua_ngan="", ten_chinh="", loai="", don_vi="", tinh_trang="",
                       so_lan="", nguon_metadata="", nguon_nhom4="",
                       bk_bbt="", bk_ma="", bk_ten="", bk_tac_gia="", bk_vi_tri="",
                       bk_diem="", bk_lan_xb="", bk_nguon_ten="")
        if rec["trang_thai"] != "ok":
            loi.append(rec)
            print("[%3d/%3d] %-8s %s" % (i, len(ds_file), rec["trang_thai"], rel[:70]))
            if rec["trang_thai"] == "thieu_nam":
                # van copy sang thu muc _KHONG_XAC_DINH_NAM voi ten goc
                if not args.dry_run:
                    os.makedirs(rec["thu_muc_dich"], exist_ok=True)
                    ten = naming.unique_path(rec["thu_muc_dich"], rec["ten_moi"], da_dung)
                    rec["ten_moi"] = ten
                    shutil.copy2(p, os.path.join(rec["thu_muc_dich"], ten))
            continue

        # dat ten duy nhat va copy
        os.makedirs(rec["thu_muc_dich"], exist_ok=True) if not args.dry_run else None
        ten = naming.unique_path(rec["thu_muc_dich"], rec["ten_moi"], da_dung)
        if ten != rec["ten_moi"]:
            rec["canh_bao"].append("ten bi trung, da them so thu tu: %s" % ten)
            rec["ten_moi"] = ten
        if not args.dry_run:
            dich = os.path.join(rec["thu_muc_dich"], rec["ten_moi"])
            if os.path.exists(dich):
                if cfg.get("ghi_de_file_dich"):
                    os.remove(dich)
                else:
                    rec["canh_bao"].append("file dich da ton tai, bo qua copy")
                    rows.append(rec)
                    continue
            if cfg.get("copy_giu_file_goc", True):
                shutil.copy2(p, dich)
            else:
                shutil.move(p, dich)
        rows.append(rec)

        if i % 10 == 0 or i == len(ds_file):
            print("  ... da xu ly %d/%d file" % (i, len(ds_file)))

        if args.verbose:
            print("[%3d] %-52s -> %s | nam=%s | N4=%s (%s) | tg=%s" % (
                i, rec["ten_goc"][:52], rec["ten_moi"][:58],
                rec["nam"], rec["don_vi"], rec.get("nguon_nhom4", "")[:34],
                rec["tac_gia"][:22]))

    # ---- bao cao
    thoi_diem = dt.datetime.now()
    if not args.dry_run:
        os.makedirs(dst, exist_ok=True)
        out = os.path.join(dst, "Bao_cao_sap_xep_%s.xlsx" % thoi_diem.strftime("%Y%m%d_%H%M%S"))
        ghi_excel(rows, loi, bo_qua, cfg, src, thoi_diem, out)
    else:
        out = "(dry-run: khong ghi bao cao)"

    print("\n================ KET QUA ================")
    print("Xu ly thanh cong      : %d file" % len(rows))
    print("Khong doc duoc nam    : %d file" % len([x for x in loi if x["trang_thai"] == "thieu_nam"]))
    print("Loi khong doc duoc    : %d file" % len([x for x in loi if x["trang_thai"] == "loi_doc"]))
    print("Bo qua (khac dinh dang): %d file" % len(bo_qua))
    bk = cfg.get("_bang_ke")
    if bk is not None:
        n_tra = len([x for x in rows if x.get("bk_diem") not in ("", None)])
        n_ma = len([x for x in rows if str(x.get("nguon_nhom4", "")).startswith("bang ke")])
        print("Bang ke sach          : tra duoc %d/%d file, lay duoc ma Nhom 4 cho %d file"
              % (n_tra, len(rows), n_ma))
        if bk.chua_map:
            print("   Ma BBT chua co trong bang ma config: %s"
                  % ", ".join("%s(%d)" % (k, v[0]) for k, v in
                              sorted(bk.chua_map.items(), key=lambda kv: -kv[1][0])[:8]))
    theo_nam = {}
    for r in rows:
        theo_nam[r["nam"]] = theo_nam.get(r["nam"], 0) + 1
    for k in sorted(theo_nam, key=lambda x: (x is None, x)):
        print("   - Nam %-6s: %d file" % (k, theo_nam[k]))
    print("Bao cao Excel         : %s" % out)
    print("Thoi gian chay        : %.1f giay" % (time.time() - t0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
