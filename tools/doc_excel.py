import os
import sys

sys.path.insert(0, r"H:\tool_files\_libs")
from openpyxl import load_workbook  # noqa: E402

path = sys.argv[1]
wb = load_workbook(path, read_only=True)
print("Cac sheet:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n=== %s (%d dong) ===" % (name, ws.max_row))
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 4:
            print("   ...")
            break
        print("   ", " | ".join("" if c is None else str(c)[:34] for c in row))
wb.close()
