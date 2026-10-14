"""Tra cuu bang ke sach da xuat ban de quy ra Nhom 4 (don vi ban hanh).

File bang ke (vi du "Tong hop SXB 2006-2025.xlsx"): moi nam la mot sheet
(2006, 2007, ... 2025). Cac cot can thiet:

    Ten sach | Tac gia | BBT (Ban bien tap) | QDXB | Lan XB | Noi in

Quy trinh:  ten sach + nam xuat ban  ->  dong tuong ung trong bang ke
            ->  cot BBT  ->  ma Nhom 4 (01..21) theo QD 68-QD/TW.

Vi du: BBT = "PL"  -> Ban sach Nha nuoc Phap luat -> 03
       BBT = "KĐLL" -> Ban sach Kinh dien Ly luan  -> 04
"""
import difflib
import os
import re
import unicodedata
import warnings

# ten cot (da bo dau, chu thuong) -> khoa noi bo
COT = {
    "ten_sach": ("ten sach",),
    "tac_gia": ("tac gia", "tac gia/dich gia", "tac gia, dich gia", "dich gia"),
    "bbt": ("bbt", "ban bien tap", "ban bt"),
    "qdxb": ("qdxb",),
    "lan_xb": ("lan xb", "lan xuat ban"),
    "noi_in": ("noi in",),
    "phuong_thuc": ("phuong thuc", "phuong thuc xb", "phuong thuc xuat ban"),
}
COT_BAT_BUOC = ("ten_sach", "bbt")


def bo_dau(s):
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("đ", "d").replace("Đ", "D")


def chuan_hoa(s):
    """Bo dau, chu thuong, bo ky tu dac biet, gop khoang trang."""
    if s is None:
        return ""
    s = bo_dau(str(s)).lower()
    s = re.sub(r"[\s_\-–—]+", " ", s)
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


TU_DUNG = {"va", "cua", "ve", "cho", "voi", "cac", "mot", "nhung", "nhu", "tai", "theo",
           "trong", "ngoai", "tren", "duoi", "den", "tu", "la", "co", "khong", "cung",
           "hoac", "bang", "thi", "ma", "nhung", "cac", "nay", "do", "khi", "de"}

# --------------------------------------------------------------------------- #
# Bang ma cung: gia tri cot "BBT" trong bang ke  ->  so Nhom 4 (QD 68-QD/TW)
#
# Danh muc don vi ban hanh:
#   01 Ban lanh dao            02 Ban sach Dang           03 Ban sach Nha nuoc Phap luat
#   04 Ban sach Kinh dien Ly luan   05 Ban sach Kinh te   06 Ban sach Quoc te
#   07 Ban sach Giao khoa Tham khao 08 Trung tam To chuc in   09 Trung tam sach Quoc gia
#   10 Chi nhanh Tp Ho Chi Minh 11 Chi nhanh Da Nang      12 Chi nhanh Can Tho
#   13 Nha in Su that          14 Vu Tai chinh - Ke toan  15 Vu Ke hoach va Thu ky bien tap
#   16 Vu To chuc can bo       17 Van phong               18 Trung tam TTNCKH
#   19 Trung tam sach Hue      20 Trung tam sach Quang Ngai  21 Trung tam sach Nha Trang
#
# Bang ke viet KHONG thong nhat giua cac nam (2007 ghi "Kinh te", tu 2008 ghi "KT"...)
# nen can liet ke du cac cach viet. Khoa: KHONG DAU, CHU THUONG.
# Muon bo sung ma moi: them vao "ban_bien_tap_bo_sung" trong config.json.
# --------------------------------------------------------------------------- #
BAN_BIEN_TAP = {
    # Ban sach Dang
    "d": "02", "dg": "02", "dang": "02", "ban sach dang": "02",
    # Ban sach Nha nuoc Phap luat
    "pl": "03", "phap luat": "03", "ban sach nha nuoc phap luat": "03",
    # Ban sach Kinh dien Ly luan
    "kdll": "04", "kd ll": "04", "ly luan": "04", "lyluan": "04", "llkd": "04", "ldll": "04",
    "kinh dien": "04", "kinh dien ly luan": "04",
    # Ban sach Kinh te
    "kt": "05", "kinh te": "05",
    # Ban sach Quoc te
    "qt": "06", "quoc te": "06",
    # Ban sach Giao khoa Tham khao
    "gk": "07", "sgk": "07", "giao khoa": "07", "giao khoa tham khao": "07",
    # Trung tam To chuc in
    "tt in": "08", "ttin": "08", "tttc in": "08", "trung tam to chuc in": "08",
    # Chi nhanh Tp Ho Chi Minh
    "hcm": "10", "tphcm": "10", "tp hcm": "10", "cn tphcm": "10", "cn tp hcm": "10",
    "cntphcm": "10",
    # Chi nhanh Da Nang
    "dn": "11", "da nang": "11", "cn da nang": "11",
    # Chi nhanh Can Tho
    "can tho": "12", "cn can tho": "12",
    # Vu Ke hoach va Thu ky bien tap
    "vu kh": "15", "vkh": "15", "vkh tkbt": "15", "vu kh tkbt": "15", "vkhtk": "15",
    "vu ke hoach": "15", "vu ke hoach va thu ky bien tap": "15",
    # Vu To chuc can bo
    "vu to chuc": "16", "vu to chuc can bo": "16",
    # Van phong
    "van phong": "17",
    # Trung tam TTNCKH
    "nckh": "18", "tt nckh": "18", "ttnckh": "18", "ttnckhxb": "18", "ban nckh": "18",
    "ban nckhxb": "18", "ban nghien cuu khxb": "18", "vu kh nckh": "18",
    "trung tam ttnckh": "18",
}


def tach_tu(s):
    return set(t for t in chuan_hoa(s).split() if t)


def _tu_chinh(tu):
    """Cac tu mang nghia (bo tu noi, bo tu qua ngan)."""
    return set(t for t in tu if len(t) >= 3 and t not in TU_DUNG)


class BangKeSach:
    """Bang ke sach xuat ban theo nam, dung de tra ma Ban bien tap (Nhom 4)."""

    def __init__(self, duong_dan, nguong=0.6, tim_nam_lan_can=1, ban_do=None):
        self.duong_dan = duong_dan
        self.nguong = float(nguong)
        self.tim_nam_lan_can = int(tim_nam_lan_can)
        # bang ma cung trong code + phan bo sung (neu co) tu config
        self.ban_do = dict(BAN_BIEN_TAP)
        for k, v in (ban_do or {}).items():
            self.ban_do[chuan_hoa(k)] = v
        self.theo_nam = {}          # '2007' -> [record, ...]
        self.chua_map = {}          # ma BBT tho -> so lan gap
        self.ten_sheet = []         # cac sheet la nam
        self.loi = ""
        self._wb = None
        self.so_dong = 0

    # ------------------------------------------------------------------ #
    # Doc file
    # ------------------------------------------------------------------ #
    def _mo(self):
        if self._wb is not None:
            return self._wb
        if not os.path.isfile(self.duong_dan):
            raise FileNotFoundError("khong thay file bang ke: %s" % self.duong_dan)
        import sys
        libs = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "_libs")
        if os.path.isdir(libs) and libs not in sys.path:
            sys.path.insert(0, libs)
        from openpyxl import load_workbook
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            self._wb = load_workbook(self.duong_dan, read_only=True, data_only=True)
        self.ten_sheet = [n for n in self._wb.sheetnames if re.fullmatch(r"(19|20)\d\d", str(n).strip())]
        if not self.ten_sheet:
            self.ten_sheet = [n for n in self._wb.sheetnames
                              if re.match(r"^(19|20)\d\d", str(n).strip())]
        return self._wb

    def _doc_sheet(self, ten_sheet):
        """Doc 1 sheet nam -> danh sach record (co cache)."""
        if ten_sheet in self.theo_nam:
            return self.theo_nam[ten_sheet]
        self._mo()
        if ten_sheet not in self._wb.sheetnames:
            self.theo_nam[ten_sheet] = []
            return []
        ws = self._wb[ten_sheet]
        ds = []
        vi_tri_cot = None
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if vi_tri_cot is None:
                if i > 12:
                    break
                vals = [chuan_hoa(c) for c in row]
                if any(v.startswith("ten sach") for v in vals):
                    vi_tri_cot = {}
                    for j, v in enumerate(vals):
                        for khoa, ten_cots in COT.items():
                            if vi_tri_cot.get(khoa) is None and any(
                                    v == t or v.startswith(t) for t in ten_cots):
                                vi_tri_cot[khoa] = j
                    if "bbt" not in vi_tri_cot:
                        break          # sheet nay khong co cot BBT (vi du 2006)
                continue
            j = vi_tri_cot.get("ten_sach")
            ten = row[j] if (j is not None and j < len(row)) else None
            ten = str(ten).strip() if ten is not None else ""
            if len(ten) < 3 or chuan_hoa(ten).startswith("ten sach"):
                continue
            rec = {"sheet": ten_sheet, "dong": i, "ten": re.sub(r"\s+", " ", ten),
                   "chuan": chuan_hoa(ten), "tu": tach_tu(ten)}
            for khoa in ("tac_gia", "bbt", "qdxb", "lan_xb", "noi_in", "phuong_thuc"):
                k = vi_tri_cot.get(khoa)
                v = row[k] if (k is not None and k < len(row)) else None
                rec[khoa] = re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""
            ds.append(rec)
        self.theo_nam[ten_sheet] = ds
        self.so_dong += len(ds)
        return ds

    def doc_tat_ca(self):
        """Doc truoc toan bo (cham hon nhung bao duoc so dong)."""
        self._mo()
        for ten in self.ten_sheet:
            self._doc_sheet(ten)
        return self.so_dong, {t: len(self.theo_nam.get(t, [])) for t in self.ten_sheet}

    # ------------------------------------------------------------------ #
    # Bang ma BBT -> Nhom 4
    # ------------------------------------------------------------------ #
    def ma_ban(self, bbt, ten_sach=""):
        """Tra ma Nhom 4 tu gia tri cot BBT. -> (ma hoac None, ghi chu)"""
        if bbt is None:
            return None, "cot BBT trong"
        raw = str(bbt).strip()
        if raw == "" or raw in ("0", "-", "nan"):
            return None, "cot BBT trong"

        ung_vien = [raw]
        # "KĐ - LL; KT" hoac "KT+Đ": tach ra tung ma
        if re.search(r"[;,+/]|\s-\s", raw):
            ung_vien += [p for p in re.split(r"[;,+/]|\s-\s", raw) if p.strip()]
        for u in ung_vien:
            for k in (chuan_hoa(u), chuan_hoa(u).replace(" ", ""), chuan_hoa(u).replace(" ", "-")):
                if k and k in self.ban_do:
                    if u != raw:
                        return self.ban_do[k], ("o BBT ghi nhieu ma (%s), da lay '%s' -> %s"
                                                % (raw, u.strip(), self.ban_do[k]))
                    return self.ban_do[k], ""
        cu = self.chua_map.get(raw) or [0, ten_sach]
        cu[0] += 1
        if not cu[1]:
            cu[1] = ten_sach
        self.chua_map[raw] = cu
        return None, "ma BBT '%s' chua co trong bang ma 'ban_bien_tap' cua config" % raw

    # ------------------------------------------------------------------ #
    # Tra cuu
    # ------------------------------------------------------------------ #
    def _diem(self, tu_q, chuoi_q, rec, tac_gia_q):
        """Diem tuong dong 0..1: chi so Dice tren tap tu + ty le chuoi (difflib).

        Chi so Dice (2*trung/(|A|+|B|)) duoc chon vi ben vung voi cac tieu de
        ngan trong bang ke (kieu 'Ve dao duc') - neu dung ty le tren tap tu ngan
        thi cac tieu de ngan de bi khop nham.
        """
        tu_r = rec["tu"]
        if not tu_r or not tu_q:
            return 0.0
        inter = len(tu_q & tu_r)
        if inter == 0:
            return 0.0
        dice = 2.0 * inter / (len(tu_q) + len(tu_r))
        if dice < 0.25:
            return dice
        seq = difflib.SequenceMatcher(None, chuoi_q, rec["chuan"]).ratio()
        diem = max(dice, seq)
        if tac_gia_q and rec["tac_gia"]:
            if len(_tu_chinh(tach_tu(tac_gia_q)) & _tu_chinh(tach_tu(rec["tac_gia"]))) >= 2:
                diem = min(1.0, diem + 0.05)
        return diem

    def tra_nhieu(self, ten_sach, nam=None, tac_gia="", so_luong=3):
        """Nhu tra() nhung tra ve danh sach ung vien tot nhat (de kiem tra tay)."""
        chuoi_q = chuan_hoa(ten_sach)
        tu_q = tach_tu(ten_sach)
        if not chuoi_q:
            return []
        ds_sheet = []
        if nam:
            ds_sheet.append(str(nam))
            for d in range(1, self.tim_nam_lan_can + 1):
                ds_sheet.append(str(int(nam) - d))
                ds_sheet.append(str(int(nam) + d))
        else:
            self._mo()
            ds_sheet = list(self.ten_sheet)
        ket = []
        for ten_sheet in ds_sheet:
            ds = self._xep_hang(ten_sheet, tu_q, chuoi_q, tac_gia)
            ket.extend(ds[:so_luong])
            if ds and ds[0][0] >= self.nguong:
                break
        ket.sort(key=lambda x: -x[0])
        ra = []
        for d, r in ket[:so_luong]:
            ra.append(dict(diem=round(d, 3), ten_sach=r["ten"], bbt=r["bbt"],
                           tac_gia=r["tac_gia"], sheet=r["sheet"], dong=r["dong"],
                           lan_xb=r["lan_xb"], noi_in=r["noi_in"]))
        return ra

    def _xep_hang(self, ten_sheet, tu_q, chuoi_q, tac_gia):
        """Xep hang ung vien trong 1 sheet. -> [(diem, rec), ...] giam dan."""
        ds = []
        for rec in self._doc_sheet(ten_sheet):
            diem = self._diem(tu_q, chuoi_q, rec, tac_gia)
            if diem > 0:
                ds.append((diem, rec))
        ds.sort(key=lambda x: -x[0])
        return ds

    def _ket_qua(self, diem, diem_2, rec, nam, ngoai_nam=False):
        ma, ghi_chu = self.ma_ban(rec["bbt"], rec["ten"])
        return {
            "bbt": rec["bbt"], "ma": ma, "ghi_chu_ma": ghi_chu,
            "ten_sach": rec["ten"], "tac_gia": rec["tac_gia"],
            "diem": round(diem, 3), "diem_2": round(diem_2, 3),
            "sheet": rec["sheet"], "dong": rec["dong"],
            "lan_xb": rec["lan_xb"], "noi_in": rec["noi_in"],
            "qdxb": rec["qdxb"], "phuong_thuc": rec["phuong_thuc"],
            "lech_nam": (int(rec["sheet"]) - int(nam)) if nam else 0,
            "ngoai_nam": ngoai_nam,
            "dat_nguong": diem >= self.nguong,
            "chac_chan": (diem >= self.nguong) and (diem - diem_2 >= 0.08),
        }

    def tra(self, ten_sach, nam=None, tac_gia="", mo_rong=False):
        """Tra sach trong bang ke theo ten sach + nam.

        Uu tien sheet DUNG NAM: neu nam do co ung vien dat nguong thi lay luon.
        Chi khi nam do khong co moi tim sang cac nam lan can (+-tim_nam_lan_can).
        mo_rong=True: neu van khong dat nguong thi tim o TAT CA cac nam khac
        (ket qua tra ve co ngoai_nam=True de canh bao).
        """
        chuoi_q = chuan_hoa(ten_sach)
        tu_q = tach_tu(ten_sach)
        if not chuoi_q or not tu_q:
            return None

        ds_sheet = []
        if nam:
            ds_sheet.append(str(nam))
            for d in range(1, self.tim_nam_lan_can + 1):
                ds_sheet.append(str(int(nam) - d))
                ds_sheet.append(str(int(nam) + d))
        else:
            self._mo()
            ds_sheet = list(self.ten_sheet)

        du_phong = None            # ung vien tot nhat du chua dat nguong
        for ten_sheet in ds_sheet:
            ds = self._xep_hang(ten_sheet, tu_q, chuoi_q, tac_gia)
            if not ds:
                continue
            diem, rec = ds[0]
            diem_2 = ds[1][0] if len(ds) > 1 else 0.0
            if diem >= self.nguong:
                return self._ket_qua(diem, diem_2, rec, nam)
            if du_phong is None or diem > du_phong[0]:
                du_phong = (diem, diem_2, rec)

        if mo_rong:
            self._mo()
            tot_nhat = None
            for ten_sheet in self.ten_sheet:
                if ten_sheet in ds_sheet:
                    continue
                ds = self._xep_hang(ten_sheet, tu_q, chuoi_q, tac_gia)
                if not ds:
                    continue
                diem, rec = ds[0]
                if diem >= self.nguong and (tot_nhat is None or diem > tot_nhat[0]):
                    tot_nhat = (diem, ds[1][0] if len(ds) > 1 else 0.0, rec)
            if tot_nhat:
                return self._ket_qua(tot_nhat[0], tot_nhat[1], tot_nhat[2], nam, ngoai_nam=True)

        if du_phong:
            return self._ket_qua(du_phong[0], du_phong[1], du_phong[2], nam)
        return None
